"""Format-specific document extractors.

Every extractor treats file content as untrusted data (rules.md §4.4): no
macros, scripts, links, or embedded objects are executed or followed.
"""

from __future__ import annotations

import csv
import io
import json
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from omnirecall.errors import EngineError, LimitExceededError, UnsupportedFormatError
from omnirecall.extraction.text import normalize_text


@dataclass
class Extracted:
    text: str
    mime: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class Format:
    kind: str
    mime: str
    label: str
    extensions: tuple[str, ...]


# The supported list shown in the UI.
DOCUMENT_FORMATS: tuple[Format, ...] = (
    Format("text", "text/plain", "Plain text", (".txt", ".text", ".log")),
    Format("markdown", "text/markdown", "Markdown", (".md", ".markdown")),
    Format("csv", "text/csv", "CSV", (".csv", ".tsv")),
    Format("json", "application/json", "JSON", (".json",)),
    Format("html", "text/html", "HTML", (".html", ".htm")),
    Format("pdf", "application/pdf", "PDF (text layer only)", (".pdf",)),
    Format(
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "Word (.docx)",
        (".docx",),
    ),
)

# Images are read with OCR (extraction/ocr.py); only the first frame is used.
IMAGE_FORMATS: tuple[Format, ...] = (
    Format("image", "image/png", "PNG image", (".png",)),
    Format("image", "image/jpeg", "JPEG image", (".jpg", ".jpeg")),
    Format("image", "image/webp", "WebP image", (".webp",)),
    Format("image", "image/bmp", "BMP image", (".bmp",)),
    Format("image", "image/tiff", "TIFF image (first page)", (".tif", ".tiff")),
    Format("image", "image/gif", "GIF image (first frame)", (".gif",)),
)

# Audio is decoded by the bundled helper and transcribed locally (D-031). Opus is
# not supported by the decoder, so .opus and WebM recordings are refused up front.
AUDIO_FORMATS: tuple[Format, ...] = (
    Format("audio", "audio/wav", "WAV audio", (".wav", ".wave")),
    Format("audio", "audio/mpeg", "MP3 audio", (".mp3",)),
    Format("audio", "audio/mp4", "M4A audio (AAC or Apple Lossless)", (".m4a",)),
    Format("audio", "audio/aac", "AAC audio", (".aac",)),
    Format("audio", "audio/flac", "FLAC audio", (".flac",)),
    Format("audio", "audio/ogg", "Ogg Vorbis audio", (".ogg", ".oga")),
    Format("audio", "audio/aiff", "AIFF audio", (".aif", ".aiff")),
    Format("audio", "audio/x-caf", "Core Audio (.caf)", (".caf",)),
)

_BY_EXTENSION = {
    ext: fmt for fmt in DOCUMENT_FORMATS + IMAGE_FORMATS + AUDIO_FORMATS for ext in fmt.extensions
}


def format_for(path: Path) -> Format:
    fmt = _BY_EXTENSION.get(path.suffix.lower())
    if fmt is None:
        raise UnsupportedFormatError(path.suffix.lower().lstrip(".") or "unknown")
    return fmt


def _decode(data: bytes) -> str:
    utf16 = data.startswith((b"\xff\xfe", b"\xfe\xff"))
    if not utf16 and b"\x00" in data[:4096]:
        raise EngineError("binary_content", "The file does not contain readable text")
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


class _HTMLText(HTMLParser):
    _SKIP = frozenset({"script", "style", "noscript", "template", "svg", "head"})
    _BLOCK = frozenset(
        {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}
    )

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "title":
            self._in_title = True
        if tag in self._SKIP:
            self._skip += 1
        if tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self._in_title = False
        if tag in self._SKIP and self._skip:
            self._skip -= 1
        if tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        if not self._skip:
            self.parts.append(data)


def _plain(data: bytes) -> tuple[str, dict[str, Any]]:
    return _decode(data), {}


def _csv(data: bytes) -> tuple[str, dict[str, Any]]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows), {"rows": len(rows)}


def _json(data: bytes) -> tuple[str, dict[str, Any]]:
    try:
        parsed = json.loads(_decode(data))
    except json.JSONDecodeError:
        raise EngineError("malformed_file", "The JSON file could not be parsed") from None
    lines: list[str] = []

    def walk(value: Any, prefix: str) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                walk(child, f"{prefix}.{key}" if prefix else str(key))
        elif isinstance(value, list):
            for child in value:
                walk(child, prefix)
        elif value is not None:
            lines.append(f"{prefix}: {value}" if prefix else str(value))

    walk(parsed, "")
    return "\n".join(lines), {}


def _html(data: bytes) -> tuple[str, dict[str, Any]]:
    parser = _HTMLText()
    parser.feed(_decode(data))
    parser.close()
    meta = {"title": normalize_text(parser.title)[:300]} if parser.title.strip() else {}
    return "".join(parser.parts), meta


def _make_pdf(max_pages: int) -> Callable[[bytes], tuple[str, dict[str, Any]]]:
    def extract(data: bytes) -> tuple[str, dict[str, Any]]:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise EngineError("encrypted_file", "Password-protected PDFs are not supported")
            pages = len(reader.pages)
            if pages > max_pages:
                raise LimitExceededError("pdf_pages", max_pages)
            texts = [(page.extract_text() or "") for page in reader.pages]
        except (PdfReadError, ValueError, KeyError, TypeError) as exc:
            raise EngineError("malformed_file", "The PDF could not be read") from exc
        text = "\n\n".join(t for t in texts if t.strip())
        meta: dict[str, Any] = {"pages": pages}
        if not text.strip():
            meta["note"] = "no_text_layer"
        return text, meta

    return extract


def _docx(data: bytes) -> tuple[str, dict[str, Any]]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise EngineError("malformed_file", "The Word document could not be read") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts), {"paragraphs": len(document.paragraphs)}


def extract_document(path: Path, max_bytes: int, max_pdf_pages: int) -> Extracted:
    fmt = format_for(path)
    size = path.stat().st_size
    if size > max_bytes:
        raise LimitExceededError("file_size", max_bytes)
    data = path.read_bytes()
    extractors: dict[str, Callable[[bytes], tuple[str, dict[str, Any]]]] = {
        "text": _plain,
        "markdown": _plain,
        "csv": _csv,
        "json": _json,
        "html": _html,
        "pdf": _make_pdf(max_pdf_pages),
        "docx": _docx,
    }
    raw, meta = extractors[fmt.kind](data)
    text = normalize_text(raw)
    if not text:
        raise EngineError("extraction_empty", "No text could be extracted from this file")
    return Extracted(text=text, mime=fmt.mime, metadata=meta | {"format": fmt.kind, "bytes": size})
